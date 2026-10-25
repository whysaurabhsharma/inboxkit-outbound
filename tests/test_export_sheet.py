import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import export
import export_sheet


def test_sheet_has_every_tab_with_bold_frozen_headers(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    import pipeline
    import run_log
    for mod in (export, run_log):
        monkeypatch.setattr(mod, "OUT", tmp_path)
    monkeypatch.setattr(run_log, "LOG", tmp_path / "run_log.csv")
    monkeypatch.setattr(export_sheet, "OUT", tmp_path)
    pipeline.run(str(ROOT / "config" / "icp.yaml"))
    path, counts, sources = export_sheet.build(tmp_path / "out.xlsx")

    wb = load_workbook(path)
    assert wb.sheetnames == ["Read me", "Leads", "Needs contact", "Held",
                             "Rejects", "Suppressed", "Run log"]
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2", ws.title
        assert all(c.font.bold for c in ws[1] if c.value), ws.title
    leads_header = [c.value for c in wb["Leads"][1]]
    assert leads_header[:9] == export.BRIEF_COLUMNS
    assert leads_header[9] == "data_source"
    assert counts["Leads"] == 39 and counts["Needs contact"] == 2
    assert (sources["live_board"], sources["synthetic_dry_run"]) == (2, 98)
    readme = " ".join(str(c.value) for row in wb["Read me"].iter_rows() for c in row)
    assert "Tier A is unreachable" in readme and "2 live_board" in readme


def test_committed_workbook_has_nothing_private():
    """The snapshot in 02-lead-list must match the clean history: no real people."""
    import re
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "02-lead-list" / "inboxkit-leads.xlsx")
    real_ok = ("belkins.io", "memoryblue.com", "jobs.lever.co/belkins/")
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(hdr, row))
            for v in (str(x) for x in row if x is not None):
                for e in re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", v):
                    assert e.endswith(".example"), (ws.title, e)
                assert "linkedin.com/in/" not in v, (ws.title, v)
                for u in re.findall(r"https?://[^\s|]+", v):
                    assert ".example" in u or any(r in u for r in real_ok), (ws.title, u)
            if rec.get("data_source") == "live_board":
                for f in ("name", "title", "email", "linkedin_url"):
                    assert rec.get(f) in (None, "Not found"), (ws.title, f, rec.get(f))
