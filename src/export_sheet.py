"""One .xlsx of the run, ready to upload to Google Sheets. Never sends.

    python run.py sheet

Reads the CSVs in data/out and writes data/out/inboxkit-leads.xlsx with
a Read me tab first, then one tab per CSV.
"""
import csv
from collections import Counter
from datetime import date
from pathlib import Path

import export

OUT = Path(__file__).resolve().parent.parent / "data" / "out"
SHEET = OUT / "inboxkit-leads.xlsx"
MAX_WIDTH = 120

# (tab, csv file, what the tab holds)
TABS = [
    ("Leads", "leads.csv",
     "Verified-valid contacts that passed every ICP check, highest score first."),
    ("Needs contact", "needs_contact.csv",
     "Accounts that passed every ICP check but have no verified contact yet."),
    ("Held", "held.csv",
     "Catch-all addresses, held back until a second check."),
    ("Rejects", "rejects.csv",
     "Duplicates, rows that failed an ICP check, and bad emails, each with a reason."),
    ("Suppressed", "suppressed.csv",
     "Customers, competitors and opt-outs, with the rule that matched."),
    ("Run log", "run_log.csv",
     "One row per pipeline run: counts, mode, credits, seconds. It holds every run "
     "since the repo was started, including development runs; the final run is the last row."),
]
LEADS_ORDER = export.BRIEF_COLUMNS + export.PROVENANCE + export.EXTRA_COLUMNS


def read(name):
    p = OUT / name
    if not p.exists():
        raise SystemExit("No {0}. Run python run.py run first.".format(p))
    with p.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return reader.fieldnames or [], list(reader)


def cell(value):
    """Whole numbers as numbers so sheets can sort them; everything else as text."""
    return int(value) if value.isdigit() and len(value) < 10 else value


def write_tab(ws, header, rows):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ws.append(header)
    for r in rows:
        ws.append([cell(r.get(h) or "") for h in header])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        longest = max([len(str(h))] + [len(str(r.get(h) or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, MAX_WIDTH)


def build(path=SHEET):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    tabs = []
    for title, name, note in TABS:
        header, rows = read(name)
        if title == "Leads":
            header = LEADS_ORDER
        tabs.append((title, note, header, rows))

    sources = Counter(r.get("data_source") for t, _, _, rows in tabs
                      if t != "Run log" for r in rows)
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        ("What this is", "inboxkit-outbound output: one ICP segment run through the pipeline. "
                         "This repo writes files and sends nothing."),
        ("Generated", date.today().isoformat()),
        ("Rows by source", "{0} live_board (real agencies from public job boards), "
                           "{1} synthetic_dry_run (invented, on .example domains). "
                           "The data_source column on every tab says which.".format(
                               sources.get("live_board", 0), sources.get("synthetic_dry_run", 0))),
        ("Tier A", "Tier A is unreachable because headcount and techstack cannot be sourced free."),
    ] + [(t, "{0} ({1} rows)".format(note, len(rows))) for t, note, _, rows in tabs]
    readme.append(["Item", "Detail"])
    for item, detail in lines:
        readme.append([item, detail])
    for c in readme[1]:
        c.font = Font(bold=True)
    readme.freeze_panes = "A2"
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    counts = {}
    for title, _, header, rows in tabs:
        write_tab(wb.create_sheet(title), header, rows)
        counts[title] = len(rows)
    wb.save(path)
    return path, counts, sources
